from io import BytesIO
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib import colors
from pptx import Presentation
from pptx.util import Inches, Pt
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from reportlab.platypus import Image


def _rows(analysis):
    rows = [["Référence", "Date", "Valeur (MMDH)"]]
    labels = {"date_courante": "T0", "trimestre_precedent": "T-1", "annee_precedente": "N-1", "arret_annuel": "Arrêté annuel"}
    for item in analysis.get("series", []):
        rows.append([labels.get(item["key"], item["key"]), item["date_arret"], item["value"]])
    return rows


def _chart_png(analysis):
    labels = {"date_courante": "T0", "trimestre_precedent": "T-1", "annee_precedente": "N-1", "arret_annuel": "Annuel"}
    series = analysis.get("series", [])
    figure, axis = plt.subplots(figsize=(7, 3.2))
    axis.bar([labels.get(item["key"], item["key"]) for item in series], [item["value"] for item in series], color="#004b87")
    axis.set_ylabel("MMDH")
    axis.set_title("Comparaison des quatre arrêtés")
    figure.tight_layout()
    buffer = BytesIO()
    figure.savefig(buffer, format="png", dpi=160)
    plt.close(figure)
    buffer.seek(0)
    return buffer


def make_pdf(title, analysis, insights):
    buffer = BytesIO()
    document = SimpleDocTemplate(buffer, pagesize=A4, title=title)
    styles = getSampleStyleSheet()
    content = [Paragraph(title, styles["Title"]), Spacer(1, 16), Paragraph("Résumé exécutif", styles["Heading2"])]
    content += [Paragraph(item, styles["BodyText"]) for item in insights]
    table = Table(_rows(analysis), repeatRows=1)
    table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#004b87")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("GRID", (0, 0), (-1, -1), .25, colors.grey), ("PADDING", (0, 0), (-1, -1), 6)]))
    chart = _chart_png(analysis)
    content += [Spacer(1, 16), Paragraph("Comparaison des quatre arrêtés", styles["Heading2"]), table, Spacer(1, 14), Image(chart, width=450, height=205)]
    document.build(content)
    return buffer.getvalue()


def make_xlsx(title, analysis, insights):
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append([title])
    summary.append([])
    summary.append(["Synthèse"])
    for insight in insights:
        summary.append([insight])
    comparison = workbook.create_sheet("Comparison")
    for row in _rows(analysis):
        comparison.append(row)
    for name in ("KPI", "T0", "T-1", "N-1", "Annual", "Top_Filiales", "Top_Sectors", "Raw_Data"):
        workbook.create_sheet(name)
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="004B87")
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def make_pptx(title, analysis, insights):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    slide.shapes.title.text = title
    slide.placeholders[1].text = f"Généré le {datetime.now().strftime('%d/%m/%Y')}"
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Résumé exécutif"
    frame = slide.placeholders[1].text_frame
    frame.text = insights[0] if insights else "Aucune analyse disponible."
    for insight in insights[1:]:
        frame.add_paragraph().text = insight
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Comparaison des quatre arrêtés"
    rows = _rows(analysis)
    table = slide.shapes.add_table(len(rows), len(rows[0]), Inches(0.6), Inches(1.4), Inches(8.8), Inches(3.5)).table
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            table.cell(r, c).text = str(value)
            for paragraph in table.cell(r, c).text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(12)
    chart = _chart_png(analysis)
    slide.shapes.add_picture(chart, Inches(1.15), Inches(5.1), width=Inches(7.5), height=Inches(3.45))
    buffer = BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()
